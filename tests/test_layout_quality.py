"""Regression tests for extraction quality issues found in real use:

* numbered headings ("4. MEANING OF MAXIM:") must stay bold headings with
  their original number — not list items renumbered from 1;
* indented/justified paragraphs must stay ONE paragraph — no lines
  drifting right, centered or split off;
* numbered lists keep their source numbering in the editor and in DOCX;
* Settings and both export actions are always reachable via the menu bar.
"""

from __future__ import annotations

from pathlib import Path

import docx

from app.document.model import BlockType, ListBlock, ParagraphBlock, TextAlignment
from app.formatting.markdown_parser import parse_markdown
from app.ocr.base import OCRLine, OCRResult, OCRWord
from app.ocr.layout import reconstruct_document


def _line(text: str, left: int, right: int, top: int, height: int = 22) -> OCRLine:
    tokens = text.split()
    width = (right - left) // max(1, len(tokens))
    words = []
    x = left
    for token in tokens:
        words.append(
            OCRWord(
                text=token, left=x, top=top, width=width - 5, height=height,
                confidence=0.9,
            )
        )
        x += width
    return OCRLine(words=words)


class TestNumberedHeading:
    def test_numbered_colon_heading_is_bold_heading_with_number(self) -> None:
        result = OCRResult(
            lines=[_line("4. MEANING OF MAXIM:", 100, 480, 100, height=26)],
            image_width=1000,
            image_height=1400,
        )
        block = reconstruct_document(result).blocks[0]
        assert isinstance(block, ParagraphBlock)
        assert block.block_type in (BlockType.HEADING, BlockType.SUBHEADING)
        assert block.spans[0].bold
        assert block.plain_text().startswith("4.")  # number preserved

    def test_lowercase_numbered_line_stays_list_item(self) -> None:
        result = OCRResult(
            lines=[
                _line("1. first ordinary item", 100, 500, 100),
                _line("2. second ordinary item", 100, 500, 130),
            ],
            image_width=1000,
            image_height=1400,
        )
        block = reconstruct_document(result).blocks[0]
        assert isinstance(block, ListBlock)
        assert block.start == 1


class TestJustifiedParagraph:
    def test_indented_first_line_does_not_split_paragraph(self) -> None:
        """The exact pattern from scanned legal text: indented first line,
        full-width middle line, short last line — must be ONE paragraph."""
        result = OCRResult(
            lines=[
                _line("A maxim is a short, well-known legal or", 430, 950, 140),
                _line(
                    "moral principle that expresses a general rule or truth,",
                    100, 950, 168,
                ),
                _line(
                    "to explain, guide, or justify decisions in law.", 100, 690, 196
                ),
            ],
            image_width=1000,
            image_height=1400,
        )
        document = reconstruct_document(result)
        paragraphs = [
            b for b in document.blocks if isinstance(b, ParagraphBlock)
        ]
        assert len(paragraphs) == 1
        assert paragraphs[0].alignment is TextAlignment.LEFT
        text = paragraphs[0].plain_text()
        assert text.startswith("A maxim")
        assert text.endswith("law.")

    def test_genuinely_centered_line_still_detected(self) -> None:
        result = OCRResult(
            lines=[_line("IN THE HIGH COURT", 380, 620, 100, height=26)],
            image_width=1000,
            image_height=1400,
        )
        block = reconstruct_document(result).blocks[0]
        assert isinstance(block, ParagraphBlock)
        assert block.alignment is TextAlignment.CENTER


class TestNumberPreservation:
    def test_markdown_list_start(self) -> None:
        document = parse_markdown("4. alpha\n5. beta")
        block = document.blocks[0]
        assert isinstance(block, ListBlock)
        assert block.start == 4

    def test_editor_renders_original_numbers(self, qapp) -> None:
        from PySide6.QtGui import QTextCursor, QTextDocument

        from app.formatting.rich_text import insert_structured_document

        qdoc = QTextDocument()
        insert_structured_document(
            QTextCursor(qdoc), parse_markdown("4. alpha\n5. beta")
        )
        numbers = []
        block = qdoc.begin()
        while block.isValid():
            text_list = block.textList()
            if text_list is not None:
                numbers.append(
                    text_list.format().start() + text_list.itemNumber(block)
                )
            block = block.next()
        assert numbers == [4, 5]

    def test_docx_export_keeps_literal_numbers(self, qapp, tmp_path: Path) -> None:
        from PySide6.QtGui import QTextCursor, QTextDocument

        from app.export import export_docx
        from app.formatting.rich_text import insert_structured_document

        qdoc = QTextDocument()
        insert_structured_document(
            QTextCursor(qdoc), parse_markdown("4. alpha\n5. beta")
        )
        output = export_docx([qdoc.toHtml()], tmp_path / "numbers.docx")
        texts = [p.text for p in docx.Document(str(output)).paragraphs]
        assert any(t.startswith("4. ") and "alpha" in t for t in texts)
        assert any(t.startswith("5. ") and "beta" in t for t in texts)


class TestMenuBarAccess:
    def test_all_key_actions_in_menus(self, qapp, monkeypatch) -> None:
        from PySide6.QtWidgets import QInputDialog, QMessageBox

        monkeypatch.setattr(
            QMessageBox, "question",
            staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes),
        )
        monkeypatch.setattr(
            QInputDialog, "getText", staticmethod(lambda *a, **k: ("", False))
        )
        from app.settings.settings_manager import SettingsManager
        from app.ui.main_window import MainWindow

        window = MainWindow(SettingsManager())
        menubar = window.menuBar()
        menu_titles = [action.text() for action in menubar.actions()]
        assert "&File" in menu_titles
        assert "&Tools" in menu_titles

        all_menu_actions: list[str] = []
        for menu_action in menubar.actions():
            menu = menu_action.menu()
            if menu is not None:
                all_menu_actions.extend(a.text() for a in menu.actions())
        assert "Settings…" in all_menu_actions
        assert "Export PDF…" in all_menu_actions
        assert "Export DOCX…" in all_menu_actions
        assert "Read All Pages" in all_menu_actions
        window.deleteLater()
