"""Tests for DOCX and PDF exporters."""

from __future__ import annotations

from pathlib import Path

import docx

from app.export import export_docx, export_pdf
from app.formatting.markdown_parser import parse_markdown

_SAMPLE_MD = """# Agreement

This is **bold** and *italic* body text.

- bullet one
- bullet two

| Item | Price |
|------|-------|
| Plot | 5,000 |
"""


def _sample_html(qapp) -> str:
    from PySide6.QtGui import QTextCursor, QTextDocument

    from app.formatting.rich_text import insert_structured_document

    qdoc = QTextDocument()
    insert_structured_document(QTextCursor(qdoc), parse_markdown(_SAMPLE_MD))
    return qdoc.toHtml()


class TestDocxExport:
    def test_structure_preserved(self, qapp, tmp_path: Path) -> None:
        html = _sample_html(qapp)
        output = export_docx([html], tmp_path / "out.docx")
        loaded = docx.Document(str(output))
        texts = [p.text for p in loaded.paragraphs]
        styles = [p.style.name for p in loaded.paragraphs]
        assert any("Agreement" in t for t in texts)
        assert "Heading 1" in styles
        assert any(s.startswith("List") for s in styles)
        assert len(loaded.tables) == 1
        assert loaded.tables[0].cell(0, 0).text == "Item"

    def test_multi_page_page_breaks(self, qapp, tmp_path: Path) -> None:
        html = _sample_html(qapp)
        output = export_docx([html, html], tmp_path / "two.docx")
        xml = docx.Document(str(output)).element.xml
        assert 'w:type="page"' in xml  # explicit page break between pages

    def test_bold_run(self, qapp, tmp_path: Path) -> None:
        html = _sample_html(qapp)
        loaded = docx.Document(str(export_docx([html], tmp_path / "b.docx")))
        bold_runs = [
            run.text
            for paragraph in loaded.paragraphs
            for run in paragraph.runs
            if run.bold
        ]
        assert any("bold" in text for text in bold_runs)


class TestPdfExport:
    def test_pdf_created(self, qapp, tmp_path: Path) -> None:
        html = _sample_html(qapp)
        output = export_pdf([html, html], tmp_path / "out.pdf")
        data = output.read_bytes()
        assert data.startswith(b"%PDF")
        assert len(data) > 1000

    def test_pdf_has_text(self, qapp, tmp_path: Path) -> None:
        import fitz

        html = _sample_html(qapp)
        output = export_pdf([html], tmp_path / "text.pdf")
        with fitz.open(output) as pdf:
            text = "".join(page.get_text() for page in pdf)
        assert "Agreement" in text
        assert "bullet one" in text
